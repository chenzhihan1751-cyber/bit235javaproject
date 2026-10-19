"""BIT235 Week 01 - assemble the Word document with code + output screenshots."""
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image

BASE = r"D:\JavaDev\JavaProjects\Week01"
IMG = os.path.join(BASE, "doc", "images")
DOCX = os.path.join(BASE, "doc", "BIT235_Week01_Tutorial_S1580887.docx")

STUDENT = "ZHIHAN CHEN"
SID = "S1580887"

EX = {
1: ("Declaring Variables",
    "Create variables with appropriate data types for your age, your height in metres, your first initial, "
    "and whether you are enrolled in BIT235. The values are hard coded (no user input) and then printed.", None),
2: ("Arithmetic Expressions",
    "Declare int a = 12; and int b = 5; and print the result of a + b, a - b, a * b, a / b and a % b. "
    "Explain why division behaves the way it does.", None),
3: ("Type Casting",
    "Declare int total = 7; and int count = 2; and calculate the average as a decimal value. "
    "Print the correct result.", None),
4: ("User Input and Comparison",
    "Ask the user to enter their age. If the age is 18 or older print \u201cYou are an adult.\u201d, "
    "otherwise print \u201cYou are under 18.\u201d", None),
5: ("Positive, Negative, or Zero",
    "Ask the user to enter a number and print one of: \u201cPositive\u201d, \u201cNegative\u201d, \u201cZero\u201d.", None),
6: ("Even or Odd",
    "Ask the user for an integer and determine whether it is even or odd (hint: use %).", None),
7: ("Simple Grading System",
    "Ask the user to enter a mark (0\u2013100) and print one of: HD (80\u2013100), D (70\u201379), C (60\u201369), "
    "P (50\u201359), F (below 50).", None),
8: ("Boolean Expressions",
    "Declare int temperature = 32; and print \u201cIt is hot.\u201d when temperature is greater than 30, "
    "\u201cIt is mild.\u201d for 20\u201330, and \u201cIt is cold.\u201d for below 20.", None),
9: ("Logical Operators",
    "Ask the user to enter a username and a password. If the username equals \u201cadmin\u201d AND the password "
    "equals \u201cjava123\u201d, print \u201cLogin successful.\u201d, otherwise print \u201cAccess denied.\u201d",
    "Why .equals() and not ==?  == compares whether the two variables point to the same String object, while "
    ".equals() compares the characters inside them. A String built at run time (such as one entered with Scanner) "
    "is a different object from the literal, so == would be false even when the text is identical and the login "
    "would fail."),
10: ("Mini Decision Program",
    "Ask the user to enter their age and whether they have a driver licence (true/false). If the age is at least 18 "
    "AND they have a licence print \u201cYou can drive independently.\u201d If the age is at least 16 but less than 18 "
    "AND they have a licence print \u201cYou can drive with restrictions.\u201d Otherwise print \u201cYou cannot drive.\u201d",
    None),
}


def add_image(doc, path, max_w=6.4, max_h=99.0):
    w, h = Image.open(path).size
    width = min(max_w, w / 96.0)
    height = width * h / w
    if height > max_h:                       # never let one screenshot get split over pages
        width = max_h * w / h
    doc.add_picture(path, width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.LEFT


doc = Document()
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(10.5)

for s in doc.sections:
    s.top_margin = Inches(0.6)
    s.bottom_margin = Inches(0.6)
    s.left_margin = Inches(0.7)
    s.right_margin = Inches(0.7)

doc.add_heading("BIT235 Object Oriented Programming", level=0)
p = doc.add_paragraph()
p.add_run("Week 1 Tutorial \u2013 Java Revision Exercises").bold = True
p2 = doc.add_paragraph()
p2.add_run(f"Name: {STUDENT}        Student ID: {SID}")
p2.runs[0].font.size = Pt(10)
p2.runs[0].font.color.rgb = RGBColor(0x33, 0x33, 0x33)

for n in range(1, 11):
    title, question, note = EX[n]
    doc.add_page_break()
    doc.add_heading(f"Exercise {n} \u2013 {title}", level=1)

    q = doc.add_paragraph()
    q.add_run("Question: ").bold = True
    q.add_run(question)

    if note:
        nb = doc.add_paragraph()
        nb.add_run(note)

    doc.add_paragraph()
    add_image(doc, os.path.join(IMG, f"ex{n}_code.png"), max_h=5.4)
    cap = doc.add_paragraph()
    cap.add_run(f"Exercise{n}.java").font.size = Pt(8.5)

    doc.add_paragraph()
    add_image(doc, os.path.join(IMG, f"ex{n}_output.png"), max_h=3.0)
    cap = doc.add_paragraph()
    cap.add_run(f"Console output \u2013 java Exercise{n}").font.size = Pt(8.5)

doc.save(DOCX)
print("saved:", DOCX, os.path.getsize(DOCX), "bytes")
