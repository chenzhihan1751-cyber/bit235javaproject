# -*- coding: utf-8 -*-
"""Fill BIT235 Workshop 2 answers (Java code + program output) into the workshop docx."""
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

SRC = r"C:\Users\陈智涵\.dsh\attachments\v1\files\c5\c5474553fde9292662c27bfc6a0299454ca96608ecf6de92909836cf3bb29b9b\Workshop 2- Fundamental Programming Structures in Java.docx"
OUT = r"D:\VSCode代码\01-学校作业\BIT235Workshop2\BIT235_Workshop2_S1580887.docx"
CODE_DIR = r"D:\VSCode代码\01-学校作业\BIT235Workshop2"
GH_URL = "https://github.com/chenzhihan1751-cyber/BIT235Workshop2"
CODE_PT = 7.0

def read(p):
    with open(p, "r", encoding="utf-8") as f:
        return f.read().replace("\r\n", "\n").rstrip("\n").split("\n")

OUTPUTS = {
 "1.1": """Enter the first integer : 42
Enter the second integer: 15
First number  = 42
Second number = 15
Difference    = 42 - 15 = 27""",
 "1.2": """Enter the temperature in Celsius: 37
37.0 degrees Celsius = 98.6 degrees Fahrenheit""",
 "1.3": """Radius   Area of the circle
---------------------------
1.00     3.14
2.50     19.63
5.00     78.50
7.00     153.86
12.00    452.16""",
 "1.4": """Principal (P)     : $5,000.00
Interest rate (r) : 4% per year
Time period (n)   : 10 years
Amount (A)        : $7,401.22
Interest earned   : $2,401.22""",
 "1.5": """a.) Volume of the Earth : 229,730,773,333.33 cubic miles
b.) Volume of the Sun   : 338,708,987,083,333,300.00 cubic miles
c.) Sun / Earth ratio   : 1,474,373.60""",
 "1.6": """Normal road consumption : 8.00 litres / 100 km
Rough road consumption  : 9.20 litres / 100 km
Distance on normal road : 500.00 km
Distance on rough road  : 434.78 km""",
 "1.7": """Original string : I love java programming
substring(7, 11): "java"
Length          : 4""",
 "1.8": """Enter the first string : Apple
Enter the second string: Apricot
Are the two given strings equal? false
First character of string 1 : A
First character of string 2 : A
Are the first characters equal? true""",
}

FILES = {
 "1.1": "Exercise1_1.java", "1.2": "Exercise1_2.java", "1.3": "Exercise1_3.java",
 "1.4": "Exercise1_4.java", "1.5": "Exercise1_5.java", "1.6": "Exercise1_6.java",
 "1.7": "Exercise1_7.java", "1.8": "Exercise1_8.java",
}

NOTE_17 = ('substring(7, 11) starts at index 7 and stops before index 11, so it returns the '
           '4-character string "java".')

doc = Document(SRC)

def new_after(anchor, text, mono=True, bold=False, size=CODE_PT,
              before_pt=0.0, after_pt=0.0, italic=False):
    """Insert a brand-new 'No Spacing' paragraph directly after `anchor`."""
    np = anchor.insert_paragraph_before("")
    anchor._p.addnext(np._p)          # move it to sit right after the anchor
    try:
        np.style = doc.styles["No Spacing"]
    except KeyError:
        pass
    pf = np.paragraph_format
    pf.space_before = Pt(before_pt)
    pf.space_after = Pt(after_pt)
    pf.line_spacing = 1.0
    pf.left_indent = Pt(0)
    pf.first_line_indent = Pt(0)
    if text == "":
        return np
    run = np.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    name = "Consolas" if mono else "Calibri"
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rPr.makeelement(qn("w:rFonts"), {})
        rPr.insert(0, rFonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(a), name)
    return np

def drop_blank_following(p, limit=12):
    el = p._p.getnext()
    removed = 0
    while el is not None and removed < limit:
        nxt = el.getnext()
        if el.tag == qn("w:p"):
            txt = "".join(t.text or "" for t in el.iter(qn("w:t")))
            if txt.strip() == "":
                el.getparent().remove(el)
                el = nxt
                removed += 1
                continue
        break

# ---- GitHub link on the cover ----
for p in doc.paragraphs:
    if p.text.strip().startswith("GitHub Link:"):
        run = p.add_run(GH_URL)
        run.font.size = Pt(9)
        run.font.name = "Consolas"
        rPr = run._element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = rPr.makeelement(qn("w:rFonts"), {})
            rPr.insert(0, rFonts)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rFonts.set(qn(a), "Consolas")
        break

# ---- collect the eight "Answer:" placeholders (they live inside table cells) ----
answers = []
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                if p.text.strip().lower() == "answer:":
                    answers.append(p)
print("found Answer: blocks =", len(answers))

order = ["1.1", "1.2", "1.3", "1.4", "1.5", "1.6", "1.7", "1.8"]
for key, ans_p in zip(order, answers):
    anchor = new_after(ans_p, "Program: " + FILES[key], mono=False, bold=True,
                       size=9.5, before_pt=2, after_pt=2)
    if key == "1.7":
        anchor = new_after(anchor, NOTE_17, mono=False, size=9, after_pt=3)
    for line in read(CODE_DIR + "\\" + FILES[key]):
        anchor = new_after(anchor, line, mono=True)
    anchor = new_after(anchor, "", mono=False, size=5)
    anchor = new_after(anchor, "Output:", mono=False, bold=True, size=9.5,
                       before_pt=4, after_pt=2)
    for line in OUTPUTS[key].split("\n"):
        anchor = new_after(anchor, line, mono=True)
    drop_blank_following(anchor)

cp = doc.core_properties
cp.author = "Chen Zhihan (S1580887)"
cp.title = "BIT235 Workshop 2 - Fundamental Programming Structures in Java"
cp.comments = "GitHub: " + GH_URL

doc.save(OUT)
print("saved:", OUT)
