# -*- coding: utf-8 -*-
"""BIT235 Workshop 3: screenshot FIRST, then the code + output text."""
import os
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

SRC = r"C:\temp\ws3.docx"
HERE = r"D:\VSCode代码\01-学校作业\BIT235-Java\JavaProjects\Week03-Workshop3"
SHOTS = r"C:\temp\ws3full"
OUT = os.path.join(HERE, "doc", "BIT235_Workshop3_S1580887.docx")
CODE_PT = 7.0
SHOT_WIDTH = Inches(6.35)

doc = Document(SRC)


def new_after(anchor, text, mono=True, bold=False, size=CODE_PT,
              before_pt=0.0, after_pt=0.0):
    np = anchor.insert_paragraph_before("")
    anchor._p.addnext(np._p)
    try:
        np.style = doc.styles["No Spacing"]
    except KeyError:
        pass
    pf = np.paragraph_format
    pf.space_before = Pt(before_pt)
    pf.space_after = Pt(after_pt)
    pf.line_spacing = 1.0
    pf.left_indent = Pt(0)
    pf.first_line_indent = Pt(0)
    if text == "":
        return np
    run = np.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    name = "Consolas" if mono else "Calibri"
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rPr.makeelement(qn("w:rFonts"), {})
        rPr.insert(0, rFonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(a), name)
    return np


def drop_blank_following(p, limit=12):
    el = p._p.getnext()
    removed = 0
    while el is not None and removed < limit:
        nxt = el.getnext()
        if el.tag == qn("w:p"):
            txt = "".join(t.text or "" for t in el.iter(qn("w:t")))
            if txt.strip() == "":
                el.getparent().remove(el)
                el = nxt
                removed += 1
                continue
        break


def read_lines(path, encoding="utf-8"):
    with open(path, "r", encoding=encoding) as f:
        return f.read().replace("\r\n", "\n").rstrip("\n").split("\n")


def code_lines(name):
    return read_lines(os.path.join(HERE, name))


def output_lines(name):
    return read_lines(os.path.join(HERE, "output", name + ".txt"), "utf-8-sig")


def emit_shot(anchor, shot_name, caption):
    anchor = new_after(anchor, caption, mono=False, bold=True, size=9,
                       before_pt=2, after_pt=3)
    np = new_after(anchor, "", mono=False, size=6)
    np.alignment = WD_ALIGN_PARAGRAPH.LEFT
    path = os.path.join(SHOTS, shot_name + ".png")
    if os.path.exists(path):
        np.add_run().add_picture(path, width=SHOT_WIDTH)
    else:
        np.add_run("[screenshot missing: " + shot_name + "]")
    return np


def emit_program(anchor, filename, outputs, header, shots=None):
    """Screenshot first (user preference), then the program text and its output."""
    if shots:
        for shot, cap in shots:
            anchor = emit_shot(anchor, shot, cap)
        anchor = new_after(anchor, "", mono=False, size=4)
    anchor = new_after(anchor, header, mono=False, bold=True, size=9.5,
                       before_pt=3, after_pt=2)
    for line in code_lines(filename):
        anchor = new_after(anchor, line, mono=True)
    for label, out_file in outputs:
        anchor = new_after(anchor, "", mono=False, size=4)
        title = "Output:" if label is None else "Output (" + label + "):"
        anchor = new_after(anchor, title, mono=False, bold=True, size=9.5,
                           before_pt=3, after_pt=2)
        for line in output_lines(out_file):
            anchor = new_after(anchor, line, mono=True)
    return anchor


QA = [
    ("Which is larger - a short or an int, and what can each hold?",
     "An int is larger: it is 32 bits wide, while a short is only 16 bits. "
     "A short holds whole numbers from -32,768 to 32,767. "
     "An int holds whole numbers from -2,147,483,648 to 2,147,483,647, "
     "so widening short -> int is always safe and is done automatically."),
    ("Which is larger - a char or a short?",
     "Neither is larger: both are 16 bits, but their ranges differ. "
     "A char is unsigned and holds 0 to 65,535 (Unicode code units), "
     "so it cannot store a negative number. A short is signed and holds "
     "-32,768 to 32,767, so it can store negatives but not values above "
     "32,767. Converting between them therefore needs an explicit cast."),
    ("Which is larger - double or a float?",
     "A double is larger and more precise: 64 bits, a range of about "
     "+/-1.7 x 10^308 and 15-16 significant decimal digits. A float is "
     "32 bits, about +/-3.4 x 10^38 and only 7 significant digits. "
     "float -> double is a safe widening conversion; double -> float can "
     "lose precision, so it needs an explicit cast."),
    ("Which is larger - an int or a float?",
     "It depends what 'larger' means. In value range a float is larger: "
     "it reaches about +/-3.4 x 10^38, while an int stops at 2,147,483,647. "
     "In precision the int wins: it stores every integer exactly up to "
     "2^31-1, whereas a float has only a 24-bit mantissa, so integers above "
     "16,777,216 are no longer stored exactly. Converting int -> float is a "
     "widening conversion in the Java language specification, but it can "
     "still lose precision for large values."),
]

anchors = []
for child in doc.element.body.iterchildren():
    if child.tag == qn("w:p"):
        txt = "".join(t.text or "" for t in child.iter(qn("w:t"))).strip()
        if txt.lower().startswith("answer"):
            anchors.append(Paragraph(child, doc))
    elif child.tag == qn("w:tbl"):
        for p_el in child.iter(qn("w:p")):
            txt = "".join(t.text or "" for t in p_el.iter(qn("w:t"))).strip()
            if txt.lower().startswith("answer"):
                anchors.append(Paragraph(p_el, doc))

print("found Answer blocks =", len(anchors))
assert len(anchors) == 12

CAP = "VS Code screenshot - the whole window, with the code in the editor and the run in the terminal:"

a = anchors[0]
emit_program(a, "WelcomeMessage.java", [(None, "WelcomeMessage")],
             "Program: WelcomeMessage.java",
             [("01_WelcomeMessage", CAP)])
drop_blank_following(a)

a = anchors[1]
emit_program(a, "PalindromeCheck.java",
             [("input: Racecar", "PalindromeCheck_racecar"),
              ("input: hello", "PalindromeCheck_hello")],
             "Program: PalindromeCheck.java",
             [("02_PalindromeCheck", CAP)])
drop_blank_following(a)

a = anchors[2]
last = a
for q, ans in QA:
    last = new_after(last, q, mono=False, bold=True, size=9.5,
                     before_pt=3, after_pt=1)
    last = new_after(last, ans, mono=False, size=9.5, after_pt=2)
drop_blank_following(a)

a = anchors[3]
emit_program(a, "NationalDebt.java", [(None, "NationalDebt")],
             "Program: NationalDebt.java",
             [("03_NationalDebt", CAP)])
drop_blank_following(a)

a = anchors[4]
emit_program(a, "CastingPrecision.java", [(None, "CastingPrecision")],
             "Program: CastingPrecision.java",
             [("04_CastingPrecision", CAP)])
drop_blank_following(a)

a = anchors[5]
emit_program(a, "AgeStringToInt.java", [(None, "AgeStringToInt")],
             "Program: AgeStringToInt.java",
             [("05_AgeStringToInt", CAP)])
drop_blank_following(a)

a = anchors[6]
emit_program(a, "BooleanToString.java", [(None, "BooleanToString")],
             "Program: BooleanToString.java",
             [("06_BooleanToString", CAP)])
drop_blank_following(a)

a = anchors[7]
last = emit_program(a, "GreetingByTime.java",
                    [("hour = 9", "GreetingByTime_day"),
                     ("hour = 20", "GreetingByTime_evening")],
                    "Program 1: GreetingByTime.java",
                    [("07_GreetingByTime", CAP)])
last = emit_program(last, "MultiplesOfFive.java", [(None, "MultiplesOfFive")],
                    "Program 2: MultiplesOfFive.java",
                    [("08_MultiplesOfFive", CAP)])
emit_program(last, "SumOneToTen.java", [(None, "SumOneToTen")],
             "Program 3: SumOneToTen.java",
             [("09_SumOneToTen", CAP)])
drop_blank_following(a)

a = anchors[8]
emit_program(a, "SmallestOfThree.java", [(None, "SmallestOfThree")],
             "Program: SmallestOfThree.java",
             [("10_SmallestOfThree", CAP)])
drop_blank_following(a)

a = anchors[9]
emit_program(a, "NumberToLetter.java",
             [("input 10", "NumberToLetter_A"),
              ("input 7", "NumberToLetter_digit")],
             "Program: NumberToLetter.java",
             [("11_NumberToLetter", CAP)])
drop_blank_following(a)

a = anchors[10]
last = emit_program(a, "SumOfSquares.java", [(None, "SumOfSquares")],
                    "Program 1 (Exercise 4A): SumOfSquares.java",
                    [("12a_SumOfSquares", CAP)])
emit_program(last, "PerfectSquares.java", [(None, "PerfectSquares")],
             "Program 2 (Exercise 4B): PerfectSquares.java",
             [("12b_PerfectSquares", CAP)])
drop_blank_following(a)

a = anchors[11]
emit_program(a, "FramingCost.java", [(None, "FramingCost")],
             "Program: FramingCost.java",
             [("13_FramingCost", CAP)])
drop_blank_following(a)

cp = doc.core_properties
cp.author = "Chen Zhihan (S1580887)"
cp.title = "BIT235 Workshop 3 - Programming Structures in Java"
cp.comments = "GitHub: https://github.com/chenzhihan1751-cyber/bit235javaproject"

os.makedirs(os.path.dirname(OUT), exist_ok=True)
doc.save(OUT)
print("saved:", OUT, os.path.getsize(OUT), "bytes")
